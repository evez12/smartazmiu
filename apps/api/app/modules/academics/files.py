"""Reads uploaded CSV/XLSX files into rows of text cells (§6.1).

Nothing is guessed silently (§13.5): a non-UTF-8 encoding is detected and reported back for confirmation,
merged cells are reported per row, and dates stay text until the column parser checks them.
"""

import csv
import io
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from enum import StrEnum

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 20_000
# An XLSX is a zip archive; this caps what a small upload may expand to in memory.
MAX_XLSX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024

_UTF8_BOM = b"\xef\xbb\xbf"
_ZIP_MAGIC = b"PK\x03\x04"


class FileEncoding(StrEnum):
    UTF8 = "utf-8"
    WINDOWS_1254 = "windows-1254"
    """Turkish Windows code page: has ç ğ ı İ ö ş ü but no ə, which Excel replaces with "?"."""
    WINDOWS_1251 = "windows-1251"
    """Cyrillic Windows code page, common in Russian-language exports."""


class UnreadableFileError(Exception):
    """The upload cannot be read at all; no import record is created."""

    def __init__(self, code: str) -> None:
        super().__init__(code)
        self.code = code


@dataclass(frozen=True, slots=True)
class FileIssue:
    code: str
    row: int | None = None
    column: str | None = None


@dataclass(frozen=True, slots=True)
class Row:
    number: int
    """Row number as shown in Excel: the header is row 1."""
    cells: dict[str, str]


@dataclass(slots=True)
class Table:
    columns: list[str]
    rows: list[Row]
    encoding: FileEncoding | None
    issues: list[FileIssue] = field(default_factory=list)


def read_table(file_name: str, data: bytes, encoding: FileEncoding | None = None) -> Table:
    if len(data) > MAX_FILE_BYTES:
        raise UnreadableFileError("import.file_too_large")
    extension = file_name.rsplit(".", 1)[-1].lower() if "." in file_name else ""
    is_zip = data.startswith(_ZIP_MAGIC)
    if extension == "xlsx" and is_zip:
        return _read_xlsx(data)
    if extension == "csv" and not is_zip:
        return _read_csv(data, encoding)
    raise UnreadableFileError("import.unsupported_format")


def detect_encoding(data: bytes) -> FileEncoding:
    if data.startswith(_UTF8_BOM):
        return FileEncoding.UTF8
    try:
        data.decode("utf-8")
    except UnicodeDecodeError:
        pass
    else:
        return FileEncoding.UTF8
    # Russian text is written almost entirely in high bytes; Azerbaijani text in cp1254 only uses them
    # for its few special letters, so the share of high bytes among letters separates the two.
    letters = [byte for byte in data if byte >= 0xC0 or chr(byte).isalpha()]
    high = sum(1 for byte in letters if byte >= 0xC0)
    return (
        FileEncoding.WINDOWS_1251
        if letters and high / len(letters) > 0.5
        else FileEncoding.WINDOWS_1254
    )


def _read_csv(data: bytes, encoding: FileEncoding | None) -> Table:
    chosen = encoding or detect_encoding(data)
    codec = "utf-8-sig" if chosen is FileEncoding.UTF8 else chosen.value
    try:
        text = data.decode(codec)
    except UnicodeDecodeError as exc:
        raise UnreadableFileError("import.encoding_mismatch") from exc

    lines = text.splitlines()
    first_line = lines[0] if lines else ""
    # Excel writes ";" in Azerbaijani and Russian locales and "," elsewhere.
    delimiter = max(",;\t", key=first_line.count)
    records = list(csv.reader(io.StringIO(text, newline=""), delimiter=delimiter))
    return _table_from_records(
        [[cell.strip() for cell in record] for record in records], chosen, issues=[]
    )


def _read_xlsx(data: bytes) -> Table:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist()) > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise UnreadableFileError("import.file_too_large")
        # Merged-cell ranges are only available outside read-only mode; the size cap above bounds memory.
        workbook = load_workbook(io.BytesIO(data), data_only=True)
    except (zipfile.BadZipFile, InvalidFileException, KeyError, ValueError, OSError) as exc:
        raise UnreadableFileError("import.unreadable") from exc

    sheet = workbook.worksheets[0]
    records = [
        [_xlsx_cell_text(value) for value in row]
        for row in sheet.iter_rows(values_only=True, max_row=MAX_ROWS + 2)
    ]
    header = [cell.strip().lower() for cell in records[0]] if records else []
    issues: list[FileIssue] = []
    for merged in sheet.merged_cells.ranges:
        column = header[merged.min_col - 1] if merged.min_col <= len(header) else None
        if merged.min_row == 1:
            issues.append(FileIssue("merged_cells", column=column))
        else:
            issues.extend(
                FileIssue("merged_cells", row=row, column=column)
                for row in range(merged.min_row, merged.max_row + 1)
            )
    return _table_from_records(records, None, issues)


def _xlsx_cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _table_from_records(
    records: list[list[str]], encoding: FileEncoding | None, issues: list[FileIssue]
) -> Table:
    while records and not any(records[-1]):
        records.pop()
    if not records:
        return Table([], [], encoding, [*issues, FileIssue("empty_file")])

    columns = [cell.strip().lower() for cell in records[0]]
    rows = [
        Row(number, {name: cell for name, cell in zip(columns, record, strict=False) if name})
        for number, record in enumerate(records[1:], start=2)
        if any(record)
    ]
    if not rows:
        issues.append(FileIssue("empty_file"))
    if len(rows) > MAX_ROWS:
        issues.append(FileIssue("too_many_rows"))
        rows = rows[:MAX_ROWS]
    return Table(columns, rows, encoding, issues)
